#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
md2docx - 마크다운 파일을 Word(.docx) 문서로 변환합니다.

한국어 문서를 염두에 두고 만들었습니다. 표가 많은 실무 문서에서
열 수가 많은 표는 자동으로 가로 페이지에 배치해 잘리지 않게 합니다.

사용법
    python3 md2docx.py 문서.md
    python3 md2docx.py 문서.md 결과.docx
    python3 md2docx.py 문서.md --font "KoPubBatang Medium" --subtitle "작성 2026-09-15"
    python3 md2docx.py 문서.md --wide-cols 6 --no-toc

필요 패키지
    pip install python-docx
    (macOS Homebrew Python이면 --break-system-packages 추가)

지원하는 문법
    제목        # ## ### ####
    표          | a | b |  (구분선 |---| 필요)
    목록        - * 1.
    인용        >
    강조        **굵게**
    코드        `인라인`
    링크        [텍스트](주소)   → 텍스트만 남기고 파란색 표시
    구분선      ---

만들어지는 문서
    A4, 본문 10pt, 줄간격 1.3, 확대비율 100% 고정
    열이 많은 표 앞뒤로 가로 방향 구역을 자동 삽입
    맨 앞에 목차 필드 삽입 (Word에서 필드 업데이트하면 채워짐)
"""
import re, os, sys, argparse

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.exit("python-docx가 필요합니다.  pip install python-docx")

INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]*\]\([^)]*\))')
LINKCOLOR = RGBColor(0x1a, 0x4d, 0x8f)
URLCOLOR = RGBColor(0x44, 0x44, 0x88)
GRAY = RGBColor(0x66, 0x66, 0x66)


class Converter:
    def __init__(self, font="맑은 고딕", mono="D2Coding", wide_cols=7,
                 body=10.0, subtitle=None, toc=True):
        self.font, self.mono, self.wide_cols = font, mono, wide_cols
        self.body, self.subtitle, self.toc = body, subtitle, toc
        self.landscape = False
        self.first_h1 = True

    # ── 서식 헬퍼
    def setfont(self, run, size, bold=False, color=None, mono=False):
        name = self.mono if mono else self.font
        run.font.name = name
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color
        run._element.rPr.rFonts.set(qn('w:eastAsia'), name)

    def add_inline(self, p, text, size, bold_all=False):
        """굵게, 인라인 코드, 링크, 벌거벗은 URL을 구분해 런으로 나눕니다."""
        for tok in INLINE.split(text):
            if not tok:
                continue
            if tok.startswith('**') and tok.endswith('**'):
                self.setfont(p.add_run(tok[2:-2]), size, True)
            elif tok.startswith('`') and tok.endswith('`'):
                self.setfont(p.add_run(tok[1:-1]), size - 0.5, bold_all, LINKCOLOR, mono=True)
            elif tok.startswith('[') and '](' in tok:
                m = re.match(r'\[([^\]]*)\]\(([^)]*)\)', tok)
                self.setfont(p.add_run(m.group(1)), size, bold_all, LINKCOLOR)
            else:
                # 긴 URL은 줄을 넘치지 않도록 작게
                for part in re.split(r'(https?://\S+)', tok):
                    if not part:
                        continue
                    if part.startswith('http'):
                        self.setfont(p.add_run(part), max(size - 2.5, 7), False, URLCOLOR)
                    else:
                        self.setfont(p.add_run(part), size, bold_all)

    @staticmethod
    def set_orient(sec, landscape):
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
            sec.left_margin = sec.right_margin = Cm(1.5)
            sec.top_margin = sec.bottom_margin = Cm(1.5)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
            sec.left_margin = sec.right_margin = Cm(1.8)
            sec.top_margin = Cm(2.0)
            sec.bottom_margin = Cm(1.8)

    @staticmethod
    def shade(cell, hexcolor):
        el = OxmlElement('w:shd')
        el.set(qn('w:fill'), hexcolor)
        cell._tc.get_or_add_tcPr().append(el)

    @staticmethod
    def fix_zoom(doc):
        """창 크기에 따라 배율이 널뛰지 않도록 100%로 고정합니다."""
        try:
            zm = doc.settings.element.find(qn('w:zoom'))
            if zm is None:
                zm = OxmlElement('w:zoom')
                doc.settings.element.insert(0, zm)
            zm.set(qn('w:val'), 'none')
            zm.set(qn('w:percent'), '100')
        except Exception:
            pass

    def add_toc(self, doc):
        p = doc.add_paragraph()
        p.add_run()
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), r'TOC \o "1-3" \h \z \u')
        inner = OxmlElement('w:r')
        t = OxmlElement('w:t')
        t.text = "여기에서 마우스 오른쪽 단추를 누르고 [필드 업데이트]를 선택하면 목차가 만들어집니다."
        inner.append(t)
        fld.append(inner)
        p._p.append(fld)

    # ── 본체
    def convert(self, src, dst):
        lines = open(src, encoding='utf8').read().split('\n')
        doc = Document()

        st = doc.styles['Normal']
        st.font.name = self.font
        st.font.size = Pt(self.body)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), self.font)
        st.paragraph_format.line_spacing = 1.3
        st.paragraph_format.space_after = Pt(2)

        hsize = {1: self.body + 5, 2: self.body + 2.5, 3: self.body + 1, 4: self.body + 0.5}
        for lvl, name in ((1, 'Heading 1'), (2, 'Heading 2'), (3, 'Heading 3')):
            s = doc.styles[name]
            s.font.name = self.font
            s.font.size = Pt(hsize[lvl])
            s.font.bold = True
            s.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
            s.element.rPr.rFonts.set(qn('w:eastAsia'), self.font)
            s.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            s.paragraph_format.space_after = Pt(5)

        self.set_orient(doc.sections[0], False)
        self.fix_zoom(doc)

        i = 0
        while i < len(lines):
            ln = lines[i]

            # 표
            if (ln.strip().startswith('|') and i + 1 < len(lines)
                    and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i + 1])):
                i = self._table(doc, lines, i)
                continue

            # 제목
            m = re.match(r'^(#{1,4})\s+(.*)$', ln)
            if m:
                lvl, txt = len(m.group(1)), m.group(2)
                p = doc.add_paragraph(style='Heading %d' % min(lvl, 3))
                for r in p.runs:
                    r.text = ''
                self.add_inline(p, txt, hsize[lvl], bold_all=True)
                if lvl == 1 and self.first_h1:
                    self.first_h1 = False
                    if self.subtitle:
                        sub = doc.add_paragraph()
                        self.setfont(sub.add_run(self.subtitle), self.body - 0.5, False, GRAY)
                    if self.toc:
                        doc.add_paragraph()
                        self.setfont(doc.add_paragraph().add_run("목  차"), self.body + 1, True)
                        self.add_toc(doc)
                        doc.add_page_break()
                i += 1
                continue

            # 목록
            m = re.match(r'^(\s*)[-*]\s+(.*)$', ln)
            if m:
                depth = len(m.group(1)) // 2
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6 + 0.5 * depth)
                p.paragraph_format.first_line_indent = Cm(-0.35)
                self.add_inline(p, "· " + m.group(2), self.body)
                i += 1
                continue
            m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', ln)
            if m:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.first_line_indent = Cm(-0.5)
                self.add_inline(p, m.group(2) + ". " + m.group(3), self.body)
                i += 1
                continue

            # 인용
            if ln.startswith('>'):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                self.add_inline(p, ln.lstrip('> ').strip(), self.body)
                i += 1
                continue

            # 구분선, 빈 줄
            if re.match(r'^\s*---+\s*$', ln):
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
                i += 1
                continue
            if ln.strip() == '':
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                self.setfont(p.add_run(''), 6)
                i += 1
                continue

            self.add_inline(doc.add_paragraph(), ln, self.body)
            i += 1

        doc.save(dst)
        return dst

    def _table(self, doc, lines, i):
        rows, j = [], i
        while j < len(lines) and lines[j].strip().startswith('|'):
            if not re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[j]):
                rows.append([c.strip() for c in lines[j].strip().strip('|').split('|')])
            j += 1
        ncol = max(len(r) for r in rows)
        rows = [(r + [''] * ncol)[:ncol] for r in rows]

        wide = ncol >= self.wide_cols
        if wide != self.landscape:
            self.set_orient(doc.add_section(WD_SECTION.NEW_PAGE), wide)
            self.landscape = wide

        fs = self.body - 2.0 if wide else self.body - 1.2
        t = doc.add_table(rows=len(rows), cols=ncol)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = t.rows[ri].cells[ci]
                cell.text = ''
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.15
                if ri == 0:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    self.shade(cell, 'EDEDED')
                self.add_inline(para, val, fs, bold_all=(ri == 0))
        doc.add_paragraph()

        if wide:
            self.set_orient(doc.add_section(WD_SECTION.NEW_PAGE), False)
            self.landscape = False
        return j


def main():
    ap = argparse.ArgumentParser(description="마크다운을 Word 문서로 변환합니다.")
    ap.add_argument("src", help="입력 .md 파일")
    ap.add_argument("dst", nargs="?", help="출력 .docx 파일 (생략하면 같은 이름)")
    ap.add_argument("--font", default="맑은 고딕", help="본문 글꼴 (기본: 맑은 고딕)")
    ap.add_argument("--mono", default="D2Coding", help="코드 글꼴 (기본: D2Coding)")
    ap.add_argument("--size", type=float, default=10.0, help="본문 크기 pt (기본: 10)")
    ap.add_argument("--wide-cols", type=int, default=7,
                    help="이 열 수 이상이면 가로 페이지에 배치 (기본: 7)")
    ap.add_argument("--subtitle", default=None, help="첫 제목 아래에 넣을 부제")
    ap.add_argument("--no-toc", action="store_true", help="목차 필드를 넣지 않음")
    a = ap.parse_args()

    dst = a.dst or os.path.splitext(a.src)[0] + ".docx"
    c = Converter(font=a.font, mono=a.mono, wide_cols=a.wide_cols,
                  body=a.size, subtitle=a.subtitle, toc=not a.no_toc)
    print("변환 완료:", c.convert(a.src, dst))


if __name__ == "__main__":
    main()
